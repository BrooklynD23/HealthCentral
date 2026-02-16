"""Tests for Excel export with formatting."""

import pytest
from datetime import datetime
from io import BytesIO
from modules.export import ExportModule


class TestExportExcel:
    """Tests for ExportModule.export_excel."""

    def setup_method(self):
        self.export = ExportModule()
        self.sample_observations = [
            {
                "analyte_canonical": "glucose",
                "value": 95.0,
                "unit": "mg/dL",
                "ref_low": 70.0,
                "ref_high": 100.0,
                "flag": None,
                "is_abnormal": False,
                "user_verified": True,
                "collected_at": datetime(2024, 1, 15),
            },
            {
                "analyte_canonical": "hemoglobin_a1c",
                "value": 6.5,
                "unit": "%",
                "ref_low": 4.0,
                "ref_high": 5.6,
                "flag": "H",
                "is_abnormal": True,
                "user_verified": False,
                "collected_at": datetime(2024, 1, 15),
            },
        ]
        self.sample_trends = [
            {"analyte": "glucose", "trend_direction": "up", "delta_percent": 8.5},
        ]

    def test_returns_valid_xlsx_bytes(self):
        result = self.export.export_excel(self.sample_observations, self.sample_trends)
        assert isinstance(result, bytes)
        # XLSX files start with PK (zip format)
        assert result[:2] == b'PK'

    def test_workbook_has_expected_sheets(self):
        from openpyxl import load_workbook
        result = self.export.export_excel(self.sample_observations, self.sample_trends)
        wb = load_workbook(BytesIO(result))
        sheet_names = wb.sheetnames
        assert "Summary" in sheet_names
        assert "Labs" in sheet_names
        assert "Medications" in sheet_names
        assert "Trends" in sheet_names

    def test_labs_sheet_has_header_row(self):
        from openpyxl import load_workbook
        result = self.export.export_excel(self.sample_observations, self.sample_trends)
        wb = load_workbook(BytesIO(result))
        ws = wb["Labs"]
        headers = [cell.value for cell in ws[1]]
        assert "Date" in headers
        assert "Analyte" in headers
        assert "Value" in headers
        assert "Flag" in headers

    def test_labs_sheet_has_data_rows(self):
        from openpyxl import load_workbook
        result = self.export.export_excel(self.sample_observations, self.sample_trends)
        wb = load_workbook(BytesIO(result))
        ws = wb["Labs"]
        assert ws.max_row == 3  # header + 2 data rows

    def test_abnormal_rows_have_fill(self):
        from openpyxl import load_workbook
        result = self.export.export_excel(self.sample_observations, self.sample_trends)
        wb = load_workbook(BytesIO(result))
        ws = wb["Labs"]
        # Row 3 is hemoglobin_a1c (abnormal)
        fill_color = ws.cell(row=3, column=1).fill.start_color.rgb
        assert fill_color is not None  # Has conditional fill

    def test_empty_observations_returns_valid_xlsx(self):
        result = self.export.export_excel([], [])
        assert isinstance(result, bytes)
        assert result[:2] == b'PK'

    def test_frozen_header_row(self):
        from openpyxl import load_workbook
        result = self.export.export_excel(self.sample_observations, self.sample_trends)
        wb = load_workbook(BytesIO(result))
        ws = wb["Labs"]
        assert ws.freeze_panes == "A2"

    def test_formula_like_strings_are_sanitized_in_excel(self):
        from openpyxl import load_workbook

        malicious_obs = [
            {
                "analyte_canonical": "\t=SUM(1,1)",
                "value": 95.0,
                "unit": " +mg/dL",
                "ref_low": 70.0,
                "ref_high": 100.0,
                "flag": "@H",
                "is_abnormal": False,
                "user_verified": False,
                "collected_at": datetime(2024, 1, 15),
            }
        ]
        malicious_trends = [
            {"analyte": "\t=TREND", "trend_direction": "\t-up", "delta_percent": 1.2},
        ]

        result = self.export.export_excel(malicious_obs, malicious_trends)
        wb = load_workbook(BytesIO(result))

        ws_labs = wb["Labs"]
        assert ws_labs.cell(row=2, column=2).value == "'\t=SUM(1,1)"
        assert ws_labs.cell(row=2, column=4).value == "' +mg/dL"
        assert ws_labs.cell(row=2, column=7).value == "'@H"

        ws_trends = wb["Trends"]
        assert ws_trends.cell(row=2, column=1).value == "'\t=TREND"
        assert ws_trends.cell(row=2, column=2).value == "'\t-up"

    def test_excel_numeric_cells_remain_numeric(self):
        from openpyxl import load_workbook

        result = self.export.export_excel(self.sample_observations, self.sample_trends)
        wb = load_workbook(BytesIO(result))
        ws = wb["Labs"]

        assert isinstance(ws.cell(row=2, column=3).value, (int, float))
        assert isinstance(ws.cell(row=2, column=5).value, (int, float))
        assert isinstance(ws.cell(row=2, column=6).value, (int, float))

    def test_summary_contains_formula_metrics(self):
        from openpyxl import load_workbook

        result = self.export.export_excel(self.sample_observations, self.sample_trends)
        wb = load_workbook(BytesIO(result))
        ws = wb["Summary"]

        assert isinstance(ws.cell(row=8, column=2).value, str)
        assert ws.cell(row=8, column=2).value.startswith("=")
        assert isinstance(ws.cell(row=9, column=2).value, str)
        assert ws.cell(row=9, column=2).value.startswith("=")
        assert isinstance(ws.cell(row=10, column=2).value, str)
        assert ws.cell(row=10, column=2).value.startswith("=")

    def test_excel_includes_data_validations(self):
        from openpyxl import load_workbook

        result = self.export.export_excel(self.sample_observations, self.sample_trends)
        wb = load_workbook(BytesIO(result))

        labs_validations = list(wb["Labs"].data_validations.dataValidation)
        meds_validations = list(wb["Medications"].data_validations.dataValidation)

        assert any(v.formula1 == '"N,L,H,LL,HH,CRITICAL"' for v in labs_validations)
        assert any(v.formula1 == '"Yes,No"' for v in labs_validations)
        assert any(v.formula1 == '"TRUE,FALSE"' for v in meds_validations)

    def test_medications_sheet_rows_when_provided(self):
        from openpyxl import load_workbook

        medications = [
            {
                "name": "Metformin",
                "generic_name": "Metformin",
                "dosage_amount": 500,
                "dosage_unit": "mg",
                "frequency": "twice_daily",
                "is_active": True,
                "reminder_enabled": True,
                "started_at": datetime(2024, 1, 1),
                "ended_at": None,
                "instructions": "Take with meals",
            }
        ]
        result = self.export.export_excel(
            self.sample_observations,
            self.sample_trends,
            medications=medications,
        )
        wb = load_workbook(BytesIO(result))
        ws = wb["Medications"]

        assert ws.max_row == 2
        assert ws.cell(row=2, column=1).value == "Metformin"
        assert ws.cell(row=2, column=3).value == "500 mg"
